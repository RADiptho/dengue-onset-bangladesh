"""Editable PowerPoint version of Figure 1, built from spec.py with native shapes only."""
import os as _os
def _repo_root():
    d = _os.path.dirname(_os.path.abspath(__file__))
    while not _os.path.exists(_os.path.join(d, '.repo_root')):
        parent = _os.path.dirname(d)
        if parent == d: raise RuntimeError('repository root (.repo_root) not found')
        d = parent
    return d + '/'
ROOT = _repo_root()

import geopandas as gpd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree
from spec import *
SW, SH = 7.2, 2.99
prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH); sl = prs.slides.add_slide(prs.slide_layouts[6])
X = lambda x: Inches(x * SW / W); Y = lambda y: Inches((H - y) * SH / H); LW = lambda d: Inches(d * SW / W); LH = lambda d: Inches(d * SH / H)
rgb = lambda h: RGBColor.from_string(h.lstrip('#'))
def text(x, y, w, h, s, size, bold=False, color='#1a1a1a', align=PP_ALIGN.CENTER):
    tb = sl.shapes.add_textbox(X(x), Y(y + h), LW(w), LH(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(s.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = 'Calibri'
def box(x, y, w, h, st):
    fc, ec = PAL[st]; s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y + h), LW(w), LH(h)); s.adjustments[0] = 0.08
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fc); s.line.color.rgb = rgb(ec); s.line.width = Pt(0.9)
    if st == 'post': s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return s
def line(pts, color, dashed=False, head=False, width=0.8):
    for i in range(len(pts) - 1):
        (a, b), (c, d) = pts[i], pts[i + 1]
        cn = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(a), Y(b), X(c), Y(d)); cn.line.color.rgb = rgb(color); cn.line.width = Pt(width)
        if dashed: cn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if head and i == len(pts) - 2:
            te = etree.SubElement(cn.line._get_or_add_ln(), qn('a:tailEnd')); te.set('type', 'triangle'); te.set('w', 'med'); te.set('len', 'med')
for k, (x, y, w, h, st, title, lines) in B.items():
    box(x, y, w, h, st)
    if k == 'valid':
        text(x, y + h - 1.6 - 2.6, w, 2.6, title, 8.0, bold=True); continue
    big = st != 'model'; nl = title.count('\n'); ty = y + h - (1.8 if big else 1.4)
    text(x, ty - 2.6 * (nl + 1), w, 2.6 * (nl + 1), title, 8.0 if big else 7.8, bold=True)
    if lines:
        off = (4.2 + 2.0 * nl) if big else 3.6
        text(x + 0.3, ty - off - 2.2 * len(lines), w - 0.6, 2.2 * len(lines), '\n'.join(lines), 6.0, color='#3a3a3a')
g = gpd.read_file(ROOT + 'data/boundaries/bgd_admin2.shp').to_crs(32646); x0, y0, x1, y1 = g.total_bounds
bx, by, bw, bh = B['data'][:4]; sc = 2.6 / (x1 - x0); ox, oy = bx + bw - 3.4, by + 0.9
for geom in g.geometry.simplify(2500):
    for p in getattr(geom, 'geoms', [geom]):
        pts = [(ox + (a - x0) * sc, oy + (b - y0) * sc) for a, b in list(p.exterior.coords)[:-1]]
        if len(pts) < 3: continue
        fb = sl.shapes.build_freeform(X(pts[0][0]), Y(pts[0][1]), scale=1.0); fb.add_line_segments([(X(a), Y(b)) for a, b in pts[1:]], close=True)
        s = fb.convert_to_shape(); s.fill.solid(); s.fill.fore_color.rgb = rgb('#C9D6E6'); s.line.color.rgb = rgb('#FFFFFF'); s.line.width = Pt(0.15)
text(bx + 1.0, by + 1.0, 7, 1.8, '64 districts', 5.6, color=PAL['pre'][1], align=PP_ALIGN.LEFT)
for pts, st, dashed in ARROWS: line(pts, PAL[st][1], dashed=bool(dashed), head=dashed is not None)
t = TL
for r, (lab, test) in enumerate(t['rows']):
    yy = t['y_top'] - r * (t['rh'] + t['gap']); lock = test == 5
    text(t['x0'] - 4.6, yy + 0.1, 4.0, t['rh'], lab, 5.8, bold=lock, color=PAL['lock'][1] if lock else '#333333', align=PP_ALIGN.RIGHT)
    for c in range(test + 1):
        col = '#C8CED6' if c < test else (PAL['lock'][1] if lock else PAL['pre'][1])
        s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, X(t['x0'] + c * t['cw'] + 0.1), Y(yy + t['rh']), LW(t['cw'] - 0.2), LH(t['rh'])); s.fill.solid(); s.fill.fore_color.rgb = rgb(col); s.line.fill.background()
    line([(t['x0'] + test * t['cw'], yy - 0.3), (t['x0'] + test * t['cw'], yy + t['rh'] + 0.3)], '#000000', width=0.9)
yb = t['y_top'] - 3 * (t['rh'] + t['gap']) - 0.6
for c, yr in enumerate(t['years']): text(t['x0'] + c * t['cw'], yb - 1.8, t['cw'], 1.8, yr, 5.2, color='#555555')
lx, ly = 1.0, 0.9
for st, lab, dashed in LEGEND:
    s = box(lx, ly, 2.6, 1.6, st); text(lx + 3.4, ly - 0.1, 0.62 * len(lab) + 3, 1.8, lab, 6.0, color='#333333', align=PP_ALIGN.LEFT); lx += 4.6 + 0.62 * len(lab)
s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, X(lx), Y(ly + 1.6), LW(2.6), LH(1.6)); s.fill.solid(); s.fill.fore_color.rgb = rgb('#C8CED6'); s.line.fill.background()
text(lx + 3.4, ly - 0.1, 10, 1.8, 'Training years', 6.0, color='#333333', align=PP_ALIGN.LEFT); lx += 4.6 + 0.62 * 14
line([(lx + 0.6, ly - 0.2), (lx + 0.6, ly + 1.8)], '#000000', width=0.9); text(lx + 1.6, ly - 0.1, 12, 1.8, 'h-week embargo', 6.0, color='#333333', align=PP_ALIGN.LEFT)
for shp in sl.shapes:
    st_ = shp._element.find(qn('p:style'))
    if st_ is not None: shp._element.remove(st_)
prs.save(ROOT + 'figures/output/Figure1_study_workflow.pptx'); print('pptx shapes', len(sl.shapes), 'pictures', sum(1 for s in sl.shapes if s.shape_type == 13))
